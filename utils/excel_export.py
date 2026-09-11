
from openpyxl import load_workbook


def export_question1(
    template_path,
    output_path,
    grid_purchase,
    total_cost,
):
    wb = load_workbook(template_path)

    ws = wb["计划购电量"]

    start_col = 2

    for i, value in enumerate(grid_purchase):
        ws.cell(
            row=2,
            column=start_col + i,
        ).value = float(value)

    ws.cell(
        row=2,
        column=start_col + len(grid_purchase),
    ).value = float(grid_purchase.sum())

    ws.cell(
        row=2,
        column=start_col + len(grid_purchase) + 1,
    ).value = float(total_cost)

    wb.save(output_path)