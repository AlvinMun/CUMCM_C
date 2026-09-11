from pathlib import Path
from openpyxl import load_workbook

from optimization.q4_optimizer import optimize_day
from utils.battery import Battery
from utils.data_loader import DataLoader
from utils.q2_plotting import plot_representative_day

PROJECT_ROOT = Path(__file__).resolve().parent
DATA_FOLDER = PROJECT_ROOT / "data"
RESULTS_FOLDER = PROJECT_ROOT / "results"
FIGURE_FOLDER = PROJECT_ROOT / "figures"

RESULTS_FOLDER.mkdir(exist_ok=True)
FIGURE_FOLDER.mkdir(exist_ok=True)

TEMPLATE_PATH = DATA_FOLDER / "result4-2.xlsx"
OUTPUT_PATH = RESULTS_FOLDER / "result4-2.xlsx"


def main():

    print("=" * 60)
    print("CUMCM 2026 国赛 C题")
    print("问题四(2)：波动电价全年优化")
    print("=" * 60)

    loader = DataLoader(DATA_FOLDER)
    attachment2 = loader.load_attachment2()
    price_data = loader.load_attachment4()
    year_data = loader.split_into_days(attachment2)

    battery = Battery(
        capacity=12000,
        initial_energy=6000,
        max_charge_power=5000,
        max_discharge_power=5000,
        charge_efficiency=0.90,
        discharge_efficiency=0.90,
    )

    representative_dates = [
        "2025-02-01",
        "2025-02-02",
        "2025-03-20",
        "2025-06-21",
        "2025-09-23",
        "2025-12-21",
    ]

    representative_names = {
        "2025-03-20": "春分",
        "2025-06-21": "夏至",
        "2025-09-23": "秋分",
        "2025-12-21": "冬至",
    }

    representative_results = {}

    wb = load_workbook(TEMPLATE_PATH)
    ws_plan = wb["计划购电量"]
    ws_charge = wb["充放电量"]
    ws_emergency = wb["紧急购电量"]

    current_soc = battery.initial_energy
    total_cost = 0
    dt = 1 / 6

    periods = [
        "0:00-4:00",
        "4:00-8:00",
        "8:00-12:00",
        "12:00-16:00",
        "16:00-20:00",
        "20:00-24:00",
    ]

    plan_row = 2
    emergency_row = 2

    for d, (day, price_day) in enumerate(zip(year_data, price_data)):

        result = optimize_day(
            price=price_day["price"],
            load_kw=day["load"],
            pv_kw=day["pv"],
            battery=battery,
            initial_energy=current_soc,
        )

        current_soc = result["end_energy"]
        total_cost += result["total_cost"]

        current_date = str(day["date"].date())

        # -------- 计划购电量（2月以后） --------
        if day["date"].month >= 2:

            ws_plan.cell(row=plan_row, column=1).value = day["date"].date()

            for t in range(144):
                ws_plan.cell(row=plan_row, column=t + 2).value = float(
                    result["planned_purchase"][t] * dt
                )

            plan_row += 1

        # -------- 充放电量（仅代表日） --------
        if current_date in representative_dates:

            charge_energy = result["charge"] * dt
            discharge_energy = result["discharge"] * dt

            base = 2 + representative_dates.index(current_date) * 6

            for k in range(6):

                row = base + k

                if k == 0:
                    ws_charge.cell(row=row, column=1).value = day["date"].date()

                ws_charge.cell(row=row, column=2).value = periods[k]

                s = k * 24
                e = s + 24

                ws_charge.cell(row=row, column=3).value = float(
                    charge_energy[s:e].sum()
                )

                ws_charge.cell(row=row, column=4).value = float(
                    discharge_energy[s:e].sum()
                )

        # -------- 紧急购电量（仅代表日） --------
        if current_date in representative_dates:

            ws_emergency.cell(row=emergency_row, column=1).value = day["date"].date()

            for t in range(144):
                ws_emergency.cell(row=emergency_row, column=t + 2).value = float(
                    result["emergency_purchase"][t] * dt
                )

            emergency_row += 1

        # -------- 典型日图 --------
        if current_date in representative_names:

            representative_results[representative_names[current_date]] = {
                "load": day["load"],
                "pv": day["pv"],
                "grid": result["planned_purchase"] * 6,
                "charge": result["charge"] * 6,
                "discharge": result["discharge"] * 6,
            }

        if (d + 1) % 50 == 0:
            print(f"已完成 {d+1}/365 天")

    wb.save(OUTPUT_PATH)

    for season, data in representative_results.items():

        plot_representative_day(
            day_name=season,
            load=data["load"],
            pv=data["pv"],
            grid=data["grid"],
            charge=data["charge"],
            discharge=data["discharge"],
            save_path=FIGURE_FOLDER / f"图4-2_{season}典型日优化调度图.png",
        )

    print("\n" + "=" * 60)
    print("问题四(2)完成！")
    print("=" * 60)
    print(f"全年总费用：{total_cost:.2f} 元")
    print(f"年末SOC：{current_soc:.2f} kWh")
    print("\n已生成：result4-2.xlsx")


if __name__ == "__main__":
    main()