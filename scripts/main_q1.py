from pathlib import Path
from openpyxl import load_workbook

from optimization.q1_optimizer import optimize_day
from utils.battery import Battery
from utils.data_loader import DataLoader
from utils.q1_plotting import plot_battery_soc

PROJECT_ROOT = Path(__file__).resolve().parent

DATA_FOLDER = PROJECT_ROOT / "data"
RESULTS_FOLDER = PROJECT_ROOT / "results"
FIGURE_FOLDER = PROJECT_ROOT / "figures"

RESULTS_FOLDER.mkdir(exist_ok=True)
FIGURE_FOLDER.mkdir(exist_ok=True)

TEMPLATE_PATH = DATA_FOLDER / "result1.xlsx"
OUTPUT_PATH = RESULTS_FOLDER / "result1.xlsx"

DT = 1 / 6  # 10 minutes = 1/6 hour


def main():

    print("=" * 60)
    print("CUMCM 2026 国赛 C题")
    print("问题一：微网最优购电调度")
    print("=" * 60)

    loader = DataLoader(DATA_FOLDER)
    df = loader.load_attachment1()

    battery = Battery(
        capacity=12000,
        initial_energy=6000,
        max_charge_power=5000,
        max_discharge_power=5000,
        charge_efficiency=0.90,
        discharge_efficiency=0.90,
    )

    price = df["电价"].to_numpy()
    load = df["小区负载"].to_numpy()
    pv = df["光伏发电预测功率"].to_numpy()

    result = optimize_day(price, load, pv, battery)

    # -----------------------------
    # Generate SOC figure
    # -----------------------------
    plot_battery_soc(
        result["battery_energy"],
        FIGURE_FOLDER / "图1_电池荷电状态变化曲线.png",
    )

    # -----------------------------
    # Export official Excel template
    # -----------------------------
    wb = load_workbook(TEMPLATE_PATH)

    ws_purchase = wb[wb.sheetnames[0]]
    ws_battery = wb[wb.sheetnames[1]]

    # Sheet 1：144个10分钟计划购电量（kWh）
    for i in range(144):
        ws_purchase.cell(row=i + 2, column=2).value = float(
            result["grid_purchase"][i]
        )

    # Total cost
    ws_purchase["D147"] = float(result["total_cost"])

    # Sheet 2：六个4小时统计
    periods = [
        "0:00-4:00",
        "4:00-8:00",
        "8:00-12:00",
        "12:00-16:00",
        "16:00-20:00",
        "20:00-24:00",
    ]

    for k in range(6):
        s = k * 24
        e = s + 24

        ws_battery.cell(row=k + 2, column=1).value = periods[k]
        ws_battery.cell(row=k + 2, column=2).value = float(
            result["charge"][s:e].sum()
        )
        ws_battery.cell(row=k + 2, column=3).value = float(
            result["discharge"][s:e].sum()
        )

    # Initial / final SOC
    ws_battery["E2"] = float(result["battery_energy"][0])
    ws_battery["E3"] = float(result["battery_energy"][-1])

    wb.save(OUTPUT_PATH)

    total_purchase = result["grid_purchase"].sum()

    print("\n" + "=" * 60)
    print("问题一完成！")
    print("=" * 60)
    print(f"全天购电成本：{result['total_cost']:.2f} 元")
    print(f"全天购电量：{total_purchase:.2f} kWh")
    print(f"初始SOC：{result['battery_energy'][0]:.2f} kWh")
    print(f"末端SOC：{result['battery_energy'][-1]:.2f} kWh")
    print("\n已生成：result1.xlsx")


if __name__ == "__main__":
    main()