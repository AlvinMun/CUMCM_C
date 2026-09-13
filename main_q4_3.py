from pathlib import Path
from openpyxl import load_workbook
import matplotlib.pyplot as plt
import numpy as np

from optimization.q4_rolling_optimizer import optimize_day_q4
from utils.battery import Battery
from utils.data_loader import DataLoader
from utils.q2_plotting import plot_representative_day

PROJECT_ROOT = Path(__file__).resolve().parent
DATA_FOLDER = PROJECT_ROOT / "data"
RESULTS_FOLDER = PROJECT_ROOT / "results"
FIGURE_FOLDER = PROJECT_ROOT / "figures"

RESULTS_FOLDER.mkdir(exist_ok=True)
FIGURE_FOLDER.mkdir(exist_ok=True)

TEMPLATE_PATH = DATA_FOLDER / "result4-3.xlsx"
OUTPUT_PATH = RESULTS_FOLDER / "result4-3.xlsx"

plt.rcParams["font.sans-serif"] = ["Microsoft YaHei"]
plt.rcParams["axes.unicode_minus"] = False



def plot_cost_comparison(q2, q3, q42, q43, save_path):

    labels = ["问题二", "问题三", "问题四(2)", "问题四(3)"]
    values = np.array([q2, q3, q42, q43]) / 10000

    plt.figure(figsize=(7, 5))
    bars = plt.bar(labels, values)

    for b, v in zip(bars, values):
        plt.text(
            b.get_x() + b.get_width() / 2,
            v + 5,
            f"{v:.1f}",
            ha="center",
            fontsize=10,
        )

    plt.ylabel("总费用(万元)")
    plt.title("图4-4 四种策略全年成本比较")
    plt.tight_layout()
    plt.savefig(save_path, dpi=300)
    plt.close()



def main():

    print("=" * 60)
    print("CUMCM 2026 国赛 C题")
    print("问题四(3)：波动电价滚动优化")
    print("=" * 60)

    loader = DataLoader(DATA_FOLDER)

    attachment2 = loader.load_attachment2()
    forecast_data = loader.load_attachment3()
    price_data = loader.load_attachment4()
    year_data = loader.split_into_days(attachment2)

    battery = Battery(
        capacity=12000,
        initial_energy=6000,
        max_charge_power=5000,
        max_discharge_power=5000,
        charge_efficiency=0.90,
        discharge_efficiency=0.90,
    )

    representative_dates = [
        "2025-02-01",
        "2025-02-02",
        "2025-03-20",
        "2025-06-21",
        "2025-09-23",
        "2025-12-21",
    ]

    representative_names = {
        "2025-03-20": "春分",
        "2025-06-21": "夏至",
        "2025-09-23": "秋分",
        "2025-12-21": "冬至",
    }

    representative_results = {}

    wb = load_workbook(TEMPLATE_PATH)

    ws_plan = wb["计划购电量"]
    ws_adjust = wb["调整购电量"]
    ws_charge = wb["充放电量"]
    ws_emergency = wb["紧急购电量"]

    current_soc = battery.initial_energy
    dt = 1 / 6

    total_cost = 0
    total_adjust_cost = 0

    periods = [
        "0:00-4:00",
        "4:00-8:00",
        "8:00-12:00",
        "12:00-16:00",
        "16:00-20:00",
        "20:00-24:00",
    ]

    plan_row = 2
    adjust_row = 2
    emergency_row = 2

    for d, (day, forecast, price_day) in enumerate(
        zip(year_data, forecast_data, price_data)
    ):

        result = optimize_day_q4(
            price=price_day["price"],
            load_kw=day["load"],
            pv_actual_kw=day["pv"],
            forecasts=forecast["forecast"],
            battery=battery,
            initial_energy=current_soc,
        )

        current_soc = result["end_energy"]

        total_cost += result["total_cost"]
        total_adjust_cost += result["adjustment_cost"]

        current_date = str(day["date"].date())


        if day["date"].month >= 2:

            ws_plan.cell(row=plan_row, column=1).value = day["date"].date()

            for t in range(144):
                ws_plan.cell(row=plan_row, column=t + 2).value = float(
                    result["planned_purchase"][t] * dt
                )

            plan_row += 1


        if day["date"].month >= 2:

            ws_adjust.cell(row=adjust_row, column=1).value = day["date"].date()

            for t in range(144):
                ws_adjust.cell(row=adjust_row, column=t + 2).value = float(
                    result["adjusted_purchase"][t] * dt
                )

            adjust_row += 1


        if current_date in representative_dates:

            charge_energy = result["charge"] * dt
            discharge_energy = result["discharge"] * dt

            base = 2 + representative_dates.index(current_date) * 6

            for k in range(6):

                row = base + k

                if k == 0:
                    ws_charge.cell(row=row, column=1).value = day["date"].date()

                ws_charge.cell(row=row, column=2).value = periods[k]

                s = k * 24
                e = s + 24

                ws_charge.cell(row=row, column=3).value = float(
                    charge_energy[s:e].sum()
                )

                ws_charge.cell(row=row, column=4).value = float(
                    discharge_energy[s:e].sum()
                )


        if current_date in representative_dates:

            ws_emergency.cell(row=emergency_row, column=1).value = day["date"].date()

            for t in range(144):
                ws_emergency.cell(row=emergency_row, column=t + 2).value = float(
                    result["emergency_purchase"][t] * dt
                )

            emergency_row += 1

        if current_date in representative_names:

            representative_results[representative_names[current_date]] = {
                "load": day["load"],
                "pv": day["pv"],
                "grid": result["adjusted_purchase"] * 6,
                "charge": result["charge"] * 6,
                "discharge": result["discharge"] * 6,
            }

        if (d + 1) % 50 == 0:
            print(f"已完成 {d+1}/365 天")

    wb.save(OUTPUT_PATH)


    for season, data in representative_results.items():

        plot_representative_day(
            day_name=season,
            load=data["load"],
            pv=data["pv"],
            grid=data["grid"],
            charge=data["charge"],
            discharge=data["discharge"],
            save_path=FIGURE_FOLDER / f"图4-3_{season}滚动优化图.png",
        )

    plot_cost_comparison(
        16979978,                # 问题二全年费用
        30243664.74,             # 问题三全年费用
        14235646.35,             # 问题四(2)全年费用
        total_cost,              # 问题四(3)全年费用
        FIGURE_FOLDER / "图4-4_四种策略全年成本比较.png",
    )

    print("\n" + "=" * 60)
    print("问题四(3)完成！")
    print("=" * 60)
    print(f"全年总费用：{total_cost:.2f} 元")
    print(f"全年调整费用：{total_adjust_cost:.2f} 元")
    print(f"年末SOC：{current_soc:.2f} kWh")
    print("\n已生成：result4-3.xlsx")


if __name__ == "__main__":
    main()